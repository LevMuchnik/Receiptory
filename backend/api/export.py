import csv
import io
import os
import re
import shutil
import zipfile
import logging
from datetime import datetime, timedelta, timezone
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse

from backend.atomic import find_original, is_contained_name
from backend.auth import require_auth
from backend.config import get_setting
from backend.currency import CONVERSION_COLUMNS, effective_conversion, reference_currency
from backend.database import get_connection
from backend.models import ExportRequest
from backend.processing.filing import download_name
from backend.storage import get_filed_path

logger = logging.getLogger(__name__)
router = APIRouter()

# The issue-date period of a document (owner, 2026-10-22). An issue date that
# is not a YYYY-MM-DD date files the document under its upload date instead:
# NULL when the model found none, but also an empty, differently shaped or
# out-of-range string (month 13), because receipt_date is model output stored
# verbatim. Without this a document the model could not date never reached
# any "Issue Date" export. submission_date is NOT NULL (migration 001). The
# query also returns it as period_date, so the zip name uses the same rule.
RECEIPT_PERIOD_SQL = (
    "CASE WHEN d.receipt_date GLOB '[0-9][0-9][0-9][0-9]-[0-9][0-9]-[0-9][0-9]*' "
    "AND substr(d.receipt_date, 6, 2) BETWEEN '01' AND '12' "
    "AND substr(d.receipt_date, 9, 2) BETWEEN '01' AND '31' "
    "THEN substr(d.receipt_date, 1, 10) ELSE substr(d.submission_date, 1, 10) END"
)

# Characters Windows refuses in a file name, and the byte cap for a stem taken
# from a sender's file name. The zip is unpacked on the accountant's machine,
# where "Invoice: Oct?.pdf" fails to extract and a name over 255 bytes fails
# everywhere (a Hebrew name gets there at ~128 characters).
_UNSAFE_NAME_CHARS = re.compile(r'[<>:"/\\|?*]')
_MAX_STEM_BYTES = 150


def _zip_safe_stem(stem: str) -> str:
    stem = _UNSAFE_NAME_CHARS.sub("_", stem).rstrip(". ")
    raw = stem.encode("utf-8")
    if len(raw) > _MAX_STEM_BYTES:
        # Cut from the middle: a filed stem ends in its hash8, which is what
        # keeps two long names apart in one folder.
        head = raw[: _MAX_STEM_BYTES - 41].decode("utf-8", "ignore")
        tail = raw[-40:].decode("utf-8", "ignore")
        stem = f"{head}-{tail}"
    return stem.rstrip(". ") or "document"


def _zip_safe_ext(ext: str) -> str:
    """The on-disk extension, as the sender chose it (".pd:f", a control
    character): anything but letters and digits becomes "_", at most 10."""
    if not ext:
        return ""
    return "." + re.sub(r"[^A-Za-z0-9]", "_", ext[1:])[:10]


def _file_candidates(row, data_dir: str):
    """Where a document's file can come from, best first: (path, name in its folder).

    The filed copy, then the original. Documents are the key (owner,
    2026-10-19): a pending, failed or unfiled document is still delivered, as
    the file that was submitted, and so is one whose filed copy is gone. The
    caller tries each in turn, so a filed copy that vanishes mid-export (a
    concurrent reprocess renames it) still leaves the original. A generator:
    the original is looked up only when the filed copy could not be used
    (find_original lists originals/ for a non-.pdf original).
    """
    stored = row["stored_filename"]
    if stored:
        # Resolved, not joined. An absolute stored_filename made os.path.join
        # return that path verbatim, os.path.exists say yes, and the export
        # copy the file into the zip as root -- with an arcname of
        # "<section>/<cat>//etc/shadow". Such a row falls back to its original.
        try:
            pdf_path = get_filed_path(stored, data_dir)
        except ValueError:
            logger.warning(
                "Document %s has a stored_filename outside filed/; exporting "
                "its original instead, if there is one", row["id"]
            )
        else:
            # isfile, not exists: a directory under that name would be
            # "written" as an empty entry and the document stamped.
            if os.path.isfile(pdf_path):
                yield pdf_path, stored, "filed"
            else:
                logger.warning(
                    "Document %s: filed copy %s is missing; exporting its "
                    "original instead, if there is one", row["id"], stored
                )
    # find_original refuses a file_hash that is not a sha256 digest, so a
    # tampered row cannot name a path outside originals/.
    original = find_original(os.path.join(data_dir, "storage"), row["file_hash"])
    if original is not None:
        # The namer behind the document page's "download original" (#69): the
        # extension is the one on disk, and a Content-Disposition value, control
        # characters or "Invoice.pdf.exe" in the sent name are cleaned out.
        ext = _zip_safe_ext(os.path.splitext(original)[1].lower())
        name = download_name(stored, row["original_filename"], row["id"], ext)
        name = _zip_safe_stem(name[: len(name) - len(ext)] if ext else name) + ext
        if not (isinstance(stored, str) and is_contained_name(stored)):
            # No filed name to borrow date and hash from: two uploads both
            # called receipt.pdf must not collide in one category folder.
            # file_hash is 64 hex here (find_original checked it).
            prefix = _slugify(f"{str(row['submission_date'])[:10]}-{row['file_hash'][:8]}")
            name = f"{prefix}-{name}"
        yield original, name, "original"


def _csv_text(value):
    """A text cell Excel would run as a formula gets a leading ' (owner,
    2026-10-22): file names and model output reach metadata.csv. Numbers are
    left alone, so a negative amount stays a number."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _slugify(value: str) -> str:
    """Collapse whitespace/unsafe filename chars into single underscores."""
    value = re.sub(r"[^\w\-]+", "_", value.strip(), flags=re.UNICODE)
    return value.strip("_")


def _export_basename(body: ExportRequest, rows: list) -> str:
    """Build a descriptive export filename: <name>_receiptory_<dates>."""
    export_name = (get_setting("export_name") or "").strip()
    if not export_name:
        names = get_setting("business_names") or []
        export_name = names[0] if names else ""
    name_part = _slugify(export_name)

    date_part = ""
    if body.preset == "month" and body.month:
        try:
            date_part = datetime.strptime(body.month, "%Y-%m").strftime("%Y_%m")
        except ValueError:
            date_part = body.month
    elif body.preset == "full_year" and body.year is not None:
        date_part = f"{body.year:04d}"
    elif body.date_from or body.date_to:
        date_part = "_to_".join(d.replace("-", "_") for d in (body.date_from, body.date_to) if d)
    else:
        # Derive range from the exported rows, else fall back to today.
        dates = sorted(str(r["period_date"])[:10].replace("-", "_") for r in rows if r["period_date"])
        if dates and dates[0] == dates[-1]:
            date_part = dates[0]
        elif dates:
            date_part = f"{dates[0]}_to_{dates[-1]}"
        else:
            date_part = datetime.now(timezone.utc).strftime("%Y_%m_%d")

    parts = [p for p in (name_part, "receiptory", date_part) if p]
    # Final slugify sanitizes any unvalidated input (e.g. date_from on the
    # since_last_export path) before it reaches the Content-Disposition header.
    return _slugify("_".join(parts))

EXPORT_CSV_FIELDS = [
    "id", "document_type", "original_filename", "stored_filename", "receipt_date",
    "vendor_name", "vendor_tax_id", "vendor_receipt_id", "client_name", "client_tax_id",
    "description", "subtotal", "tax_amount", "total_amount", "currency",
    "payment_method", "payment_identifier", "language", "status",
    "submission_date", "submission_channel", "category_name",
    # Appended (#72), so a sheet reading earlier columns by position keeps working.
    "converted_amount", "conversion_rate", "conversion_rate_date", "conversion_ref",
    # Appended: the file's path inside the zip, empty when no file could be
    # added. An original has no stored_filename to match a row on.
    "exported_file",
]


@router.post("/export")
def export_documents(body: ExportRequest, request: Request, username: str = Depends(require_auth)):
    data_dir = request.app.state.data_dir

    # Build query conditions. Only deletion keeps a document out: no status
    # filter (owner, 2026-10-19: everything submitted in the period is
    # exported; a pending, failed or needs_review document goes out as its
    # original). What reaches Receiptory at all is decided at ingestion (the
    # Gmail labels, the Telegram chat, an upload); an explicit body.status
    # below still narrows an API call.
    conditions = ["d.is_deleted = 0"]
    params: list = []

    # If specific document IDs provided, filter to those only
    if body.document_ids:
        placeholders = ",".join("?" * len(body.document_ids))
        conditions.append(f"d.id IN ({placeholders})")
        params.extend(body.document_ids)

    date_col = "d.submission_date" if body.date_basis == "ingestion" else RECEIPT_PERIOD_SQL

    if not body.document_ids and body.preset == "since_last_export":
        conditions.append("d.last_exported_date IS NULL")
    elif body.preset == "month" and body.month:
        try:
            month_dt = datetime.strptime(body.month, "%Y-%m")
        except ValueError:
            raise HTTPException(status_code=422, detail="month must be in YYYY-MM format")
        if month_dt.month == 12:
            next_month_dt = month_dt.replace(year=month_dt.year + 1, month=1)
        else:
            next_month_dt = month_dt.replace(month=month_dt.month + 1)
        conditions.append(f"{date_col} >= ?")
        conditions.append(f"{date_col} < ?")
        params.extend([month_dt.strftime("%Y-%m-01"), next_month_dt.strftime("%Y-%m-01")])
    elif body.preset == "full_year" and body.year is not None:
        conditions.append(f"{date_col} >= ?")
        conditions.append(f"{date_col} < ?")
        params.extend([f"{body.year:04d}-01-01", f"{body.year + 1:04d}-01-01"])
    else:
        if body.date_from:
            try:
                date_from_dt = datetime.strptime(body.date_from, "%Y-%m-%d")
            except ValueError:
                raise HTTPException(status_code=422, detail="date_from must be in YYYY-MM-DD format")
            conditions.append(f"{date_col} >= ?")
            params.append(date_from_dt.strftime("%Y-%m-%d"))
        if body.date_to:
            try:
                dt = datetime.strptime(body.date_to, "%Y-%m-%d")
            except ValueError:
                raise HTTPException(status_code=422, detail="date_to must be in YYYY-MM-DD format")
            if body.date_basis == "ingestion":
                # submission_date is a full ISO timestamp; use exclusive upper bound
                # to include all documents ingested up to end-of-day UTC on date_to
                next_day = (dt + timedelta(days=1)).strftime("%Y-%m-%d")
                conditions.append(f"{date_col} < ?")
                params.append(next_day)
            else:
                conditions.append(f"{date_col} <= ?")
                params.append(dt.strftime("%Y-%m-%d"))

    if body.status:
        conditions.append("d.status = ?")
        params.append(body.status)
    if body.category_id:
        conditions.append("d.category_id = ?")
        params.append(body.category_id)
    if body.document_type:
        conditions.append("d.document_type = ?")
        params.append(body.document_type)
    if body.section:
        conditions.append("c.section = ?")
        params.append(body.section)

    where = " AND ".join(conditions)

    with get_connection() as conn:
        rows = conn.execute(
            f"""SELECT d.*, c.name as category_name, c.section as category_section,
                       {date_col} AS period_date
                FROM documents d
                LEFT JOIN categories c ON d.category_id = c.id
                WHERE {where}
                ORDER BY {date_col}""",
            params,
        ).fetchall()

    # Documents with no file in the zip (neither copy on disk). They must NOT
    # be stamped as exported below: preset="since_last_export" filters on
    # last_exported_date IS NULL, so marking an undelivered document would
    # retire it from every future incremental export, silently and forever.
    # Their count goes back in X-Export-Omitted so the page can say so.
    omitted: set[int] = set()
    # Document id -> its path inside the zip (the exported_file column).
    exported_file: dict[int, str] = {}
    # Delivered as the original before processing finished (pending, failed,
    # unfiled needs_review): blank spreadsheet fields and, stamped, they will
    # not come again. The page says how many (owner, 2026-10-22).
    unprocessed = 0
    # metadata.xlsx could not be written; metadata.csv has the same rows.
    excel_failed = False

    # Build zip in memory
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        # Add PDFs organized by section/category
        for row in rows:
            # Slugified, because BOTH of these reach the zip arcname and neither
            # is validated on the way in: categories.name is a bare `str` in
            # CategoryCreate, and section is only a Literal on write, not on
            # read. zipfile.ZipInfo.from_file normpaths the arcname and strips a
            # LEADING separator but keeps ".." intact -- verified:
            # "expense/../../../tmp/evil/x.pdf" is stored as
            # "../../tmp/evil/x.pdf", which unzip happily writes outside the
            # extraction directory. Hardening only the filename would leave two
            # of three path components open, which is the exact mistake #45 cost
            # us on scanner_test_frames.frame_path.
            section = _slugify(row["category_section"] or "") or "other"
            # Prefixed with the category id, because _slugify is many-to-one:
            # "Travel/Meals", "Travel Meals" and "Travel_Meals" all collapse to
            # one folder, and a name made entirely of punctuation slugs to ""
            # and would land in the real "uncategorized" folder. Nothing is
            # lost either way (filenames carry a hash suffix) but documents
            # would be filed under the wrong heading in the archive someone
            # else reads. Ids start at 1, so 0 is a fallback no row can claim.
            cat_id = row["category_id"] or 0
            cat_name = f"{cat_id}-" + (_slugify(row["category_name"] or "") or "uncategorized")
            for path, name, kind in _file_candidates(row, data_dir):
                arcname = f"{section}/{cat_name}/{name}"
                # Opened before the zip entry is started, so a copy that cannot
                # be read moves on to the next candidate with nothing written:
                # a filed copy unlinked by a concurrent reprocess (export takes
                # no lock), or a permission or I/O error on the NAS share. One
                # bad file must not cost the whole archive. An error during the
                # copy itself still fails the export: the entry would be torn.
                src = None
                try:
                    # O_NOFOLLOW: get_filed_path and find_original refused a
                    # symlink a moment ago; this keeps it refused at open.
                    src = os.fdopen(os.open(path, os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)), "rb")
                    # strict_timestamps=False: an original keeps its sender's
                    # mtime (a scanner with a dead clock writes 1970), which
                    # zip cannot store; it is clamped to 1980 instead of
                    # failing every export that includes the document.
                    zinfo = zipfile.ZipInfo.from_file(path, arcname, strict_timestamps=False)
                except OSError as e:
                    if src is not None:
                        src.close()
                    logger.warning(
                        "Document %s: %s could not be read (%s); trying the next copy",
                        row["id"], name, e.__class__.__name__
                    )
                    continue
                zinfo.compress_type = zf.compression
                with src, zf.open(zinfo, "w") as dest:
                    shutil.copyfileobj(src, dest, 1024 * 8)
                exported_file[row["id"]] = arcname
                if kind == "original" and row["status"] != "processed":
                    unprocessed += 1
                break
            else:
                omitted.add(row["id"])
                logger.warning(
                    "Document %s has neither a filed copy nor an original on "
                    "disk; only its metadata row is exported", row["id"]
                )

        # A conversion computed for another reference currency is stale (#72):
        # export the document as unconverted rather than mix currencies.
        ref = reference_currency()
        conversions = {row["id"]: effective_conversion(row, ref) for row in rows}

        def _cell(row, f):
            if f in CONVERSION_COLUMNS:
                return conversions[row["id"]][f]
            if f == "exported_file":
                return exported_file.get(row["id"], "")
            return row[f] if f in row.keys() else None

        # Add CSV metadata
        csv_buf = io.StringIO()
        writer = csv.DictWriter(csv_buf, fieldnames=EXPORT_CSV_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({f: _csv_text(_cell(row, f)) for f in EXPORT_CSV_FIELDS})
        zf.writestr("metadata.csv", csv_buf.getvalue())

        # Add Excel metadata
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Documents"
            ws.append(EXPORT_CSV_FIELDS)
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

            def _xlsx(v):
                # openpyxl raises on control characters, and names and text in
                # these cells come from senders and the model.
                return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

            for row in rows:
                ws.append([_xlsx(_cell(row, f)) for f in EXPORT_CSV_FIELDS])
            # Text, never formulas: openpyxl stores a string starting with "="
            # as one, and file names and model output reach these cells.
            for line in ws.iter_rows(min_row=2):
                for cell in line:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            # Auto-size columns
            for col in ws.columns:
                max_len = max((len(str(cell.value or "")) for cell in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
            xlsx_buf = io.BytesIO()
            wb.save(xlsx_buf)
            zf.writestr("metadata.xlsx", xlsx_buf.getvalue())
        except Exception as e:
            excel_failed = True
            logger.warning(f"Excel export failed, skipping: {e}")

    # Update last_exported_date
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    # Every delivered document is stamped, converted or not, processed or
    # not (owner, 2026-10-19 and 2026-10-22): the accountant works from the
    # documents; a conversion or a pending status is never a reason to hold
    # one back or send it twice.
    doc_ids = [row["id"] for row in rows if row["id"] not in omitted]
    if omitted:
        logger.warning(
            "%d document(s) had no file in this export and were left unstamped, "
            "so they stay eligible for the next since_last_export run",
            len(omitted),
        )
    if doc_ids:
        placeholders = ",".join("?" * len(doc_ids))
        with get_connection() as conn:
            conn.execute(
                # The first delivery is kept (owner, 2026-10-22): re-exporting a
                # month does not move it. since_last_export tests only NULL.
                f"UPDATE documents SET last_exported_date = COALESCE(last_exported_date, ?) "
                f"WHERE id IN ({placeholders})",
                [now] + doc_ids,
            )

    buf.seek(0)
    filename = f"{_export_basename(body, rows)}.zip"
    # ASCII fallback for legacy clients; RFC 5987 for unicode (e.g. Hebrew names).
    ascii_name = filename.encode("ascii", "ignore").decode("ascii") or "receiptory_export.zip"
    disposition = (
        f"attachment; filename=\"{ascii_name}\"; "
        f"filename*=UTF-8''{quote(filename)}"
    )
    return StreamingResponse(
        buf,
        media_type="application/zip",
        headers={
            "Content-Disposition": disposition,
            "X-Export-Omitted": str(len(omitted)),
            "X-Export-Unprocessed": str(unprocessed),
            "X-Export-Excel-Failed": "1" if excel_failed else "0",
        },
    )
